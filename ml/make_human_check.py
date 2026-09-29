"""Build the blind labelling workbook for a human check of the wet/dry gold labels.

The gold labels the model is trained and tested on were assigned by Claude,
reading each record. This workbook lets one or more people who know Emory's
labs label a random sample of the locked test set themselves, under the same
rubric, without seeing Claude's label or the model's. `human_check.py` then
scores the filled workbooks: how often the people agree with Claude's labels,
with each other, and with the model.

  python ml/make_human_check.py Emory_Lab_Atlas_v85.html --n 150 --out ml/human_check

Writes labelling_sheet.xlsx (give this to the labellers) and key.csv (keep it:
it maps the sheet's codes back to atlas records, and is the only link between
them). The sample and its order are fixed by the seed, so re-running gives the
same workbook.
"""
import argparse
import csv
import json
import os
import random
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import wetdry as W  # noqa: E402

SEED = 20260926
FONT = 'Arial'
INPUT = PatternFill('solid', fgColor='FFF2CC')      # the only cells a labeller edits
HEAD = PatternFill('solid', fgColor='1F3A4D')
BAND = PatternFill('solid', fgColor='EEF4F8')
GREY = PatternFill('solid', fgColor='F2F2F2')
THIN = Side(style='thin', color='C9D6DF')
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
MAX_TITLES = 60
LABELS = ['W', 'D', 'H', 'U']


def f(size=10, bold=False, italic=False, color='000000', underline=None):
    return Font(name=FONT, size=size, bold=bold, italic=italic, color=color, underline=underline)


def year_of(e):
    m = re.search(r'(19|20)\d\d', str(e[1] if isinstance(e, list) and len(e) > 1 else ''))
    return int(m.group(0)) if m else 0


def record(p, D):
    """What the labeller reads: the same fields Claude read, nothing derived."""
    pubs, seen = [], set()
    for e in sorted(p.get('p') or [], key=year_of, reverse=True):
        raw = str(e[0] if isinstance(e, list) else e or '').strip()
        key = W.clean_title(raw)
        if len(key) >= 12 and key not in seen:
            seen.add(key)
            y = year_of(e)
            pubs.append((raw, y or None))
    units = ['%s (%s)' % (D['depts'][u]['n'], D['insts'][D['depts'][u]['i']]['k']) for u in (p.get('d') or [])]
    return {
        'name': p['n'], 'email': p.get('e') or '', 'degrees': p.get('g') or '',
        'units': units, 'profile': W.bio_of(p).strip(), 'grants': list(p.get('gn') or []),
        'pubs': pubs, 'url': p.get('u') or '',
    }


def build(D, people, out):
    wb = Workbook()

    # ------------------------------------------------------------ instructions
    ws = wb.active
    ws.title = 'Instructions'
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 3
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 100
    r = 2
    ws.cell(r, 2, 'Emory Lab Atlas: wet/dry labelling check').font = f(16, True, color='1F3A4D')
    r += 1
    ws.cell(r, 2, '%d investigators, one row each on the Label sheet. Allow about 2 minutes a person.' % len(people)).font = f(10, italic=True, color='546C78')
    r += 2
    blocks = [
        ('Why', 'The atlas places every investigator on a wet/dry axis with a model. Its accuracy was measured against '
                'labels that Claude assigned by reading each record. This sheet checks those labels: you label the same '
                'people yourself, without seeing Claude’s label or the model’s, and the two are compared.'),
        ('Keep it blind', 'Do not open the atlas, the repository or anyone else’s copy of this sheet while you label. '
                          'If two of you are labelling, work separately and do not compare notes until both are done.'),
        ('What to do', '1. Go to the Label sheet. For each row, click "Read record" to see the person’s full record '
                       '(publication titles, profile text, grants, units, degrees), then click "Back to label".\n'
                       '2. In the yellow Label cell pick W, D, H or U from the list (definitions below).\n'
                       '3. In the yellow Sure? cell pick 1 if you are sure, 2 if it is a judgment call.\n'
                       '4. Optionally, a short note: the reason, especially for H, U or anything surprising.\n'
                       '5. Optional, and only after step 2: if you know the lab or look at its website and would label it '
                       'differently from what the record alone suggests, put that label in "After web check".\n'
                       'Edit only the yellow cells. Row 5 on the Label sheet is a worked example and is not scored.'),
        ('W  wet', 'Most of the output rests on bench or laboratory experiments that the group itself runs: cells, '
                   'molecules, proteins, nucleic acids, microbes, tissues and specimens processed in a lab; animal models '
                   '(mouse, rat, non-human primate, zebrafish, fly, worm); chemistry and synthesis; structural biology; '
                   'biomaterials and devices built and tested at the bench; electrophysiology in tissue or animals; '
                   'experimental imaging of specimens or animals.'),
        ('D  dry', 'Most of the output is research without a bench: computational biology, bioinformatics, statistics; '
                   'epidemiology and population health; clinical research on patients, records, registries or trials; '
                   'case reports; surveys, interviews, qualitative and behavioural research; health services, policy, '
                   'economics and education; image interpretation. Clinical research counts as dry: the axis is bench '
                   'versus no bench, not laboratory versus computer.'),
        ('H  hybrid', 'Both modes are a substantial, deliberate part of the output (roughly a quarter or more each), or '
                      'the group pairs bench experiments with computational or clinical-cohort analysis as its method '
                      '(mouse models alongside patient cohorts; sequencing experiments alongside the bioinformatics '
                      'built to read them).'),
        ('U  can’t tell', 'Too little to judge: no profile and only a couple of uninformative titles, or only '
                               'editorials, commentaries and guidelines.'),
        ('Reading rules', '• Judge the person, not the department. A pathologist can run a bench lab or only read slides.\n'
                          '• Weigh the bulk of the titles, not the most striking one. One mouse paper among forty '
                          'clinical papers is still dry.\n'
                          '• Multi-centre trials and consortium papers count as what they are (a clinician’s '
                          'trial papers are dry).\n'
                          '• Set aside titles plainly written by a namesake (a different field or era).\n'
                          '• Clinic addresses, board certifications and languages spoken say someone is a clinician; '
                          'they do not by themselves say whether they also run a bench.'),
        ('When done', 'Save the file with your initials in the name (for example labelling_sheet_JS.xlsx) and send it back. '
                      'The Progress sheet shows how many rows are left.'),
    ]
    for head, body in blocks:
        ws.cell(r, 2, head).font = f(11, True, color='1F3A4D')
        c = ws.cell(r, 3, body)
        c.font = f(10)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.cell(r, 2).alignment = Alignment(vertical='top')
        ws.row_dimensions[r].height = 14 * max(1, sum(len(line) // 140 + 1 for line in body.split('\n')))
        r += 2

    # ------------------------------------------------------------- label sheet
    lab = wb.create_sheet('Label')
    rec = wb.create_sheet('Records')
    cols = [('Code', 8), ('Name', 26), ('Units', 38), ('Degrees', 12), ('Titles', 8), ('Record', 13),
            ('Label', 9), ('Sure?', 8), ('Note', 40), ('After web check', 14)]
    for j, (_, w) in enumerate(cols, 1):
        lab.column_dimensions[chr(64 + j)].width = w
    lab['A1'] = 'Label each person W, D, H or U. Edit only the yellow cells.'
    lab['A1'].font = f(12, True, color='1F3A4D')
    lab['A2'] = ('W wet (bench)   D dry (no bench; clinical counts as dry)   H hybrid (both, a quarter or more each)   '
                 'U can’t tell.   Sure?: 1 sure, 2 judgment call.   Full definitions on the Instructions sheet.')
    lab['A2'].font = f(9, italic=True, color='546C78')
    hr = 4
    for j, (h, _) in enumerate(cols, 1):
        c = lab.cell(hr, j, h)
        c.font = f(10, True, color='FFFFFF')
        c.fill = HEAD
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = BOX
    lab.row_dimensions[hr].height = 22

    # worked example (fictional, not scored)
    ex = hr + 1
    for j, v in enumerate(['EX', '(example) A. Researcher', 'Pharmacology & Chemical Biology (SOM)', 'PhD', 34,
                           '(example)', 'W', 1, 'Most titles are mouse and cell-line experiments', ''], 1):
        c = lab.cell(ex, j, v)
        c.font = f(9, italic=True, color='7F7F7F')
        c.fill = GREY
        c.border = BOX
        c.alignment = Alignment(vertical='top', wrap_text=True)

    dv_lab = DataValidation(type='list', formula1='"W,D,H,U"', allow_blank=True, showErrorMessage=True,
                            errorTitle='Label', error='Pick W, D, H or U.')
    dv_web = DataValidation(type='list', formula1='"W,D,H,U"', allow_blank=True, showErrorMessage=True,
                            errorTitle='After web check', error='Pick W, D, H or U, or leave blank.')
    dv_conf = DataValidation(type='list', formula1='"1,2"', allow_blank=True, showErrorMessage=True,
                             errorTitle='Sure?', error='Pick 1 (sure) or 2 (judgment call).')
    for dv in (dv_lab, dv_web, dv_conf):
        lab.add_data_validation(dv)

    # records sheet layout
    rec.sheet_view.showGridLines = False
    rec.column_dimensions['A'].width = 16
    rec.column_dimensions['B'].width = 110
    rr = 1
    first = ex + 1
    for n, (code, i) in enumerate(people):
        row = first + n
        x = record(D['pis'][i], D)
        start = rr
        # ---- record block
        c = rec.cell(rr, 1, code)
        c.font = f(12, True, color='FFFFFF')
        c.fill = HEAD
        c = rec.cell(rr, 2, x['name'] + ('   ·   ' + x['degrees'] if x['degrees'] else ''))
        c.font = f(12, True, color='FFFFFF')
        c.fill = HEAD
        rr += 1
        # HYPERLINK formulas with a '#' target stay inside the file, so a labeller
        # can save their copy under any name and the links still work
        c = rec.cell(rr, 2, '=HYPERLINK("#Label!G%d","← Back to label (row %d)")' % (row, row))
        c.font = f(10, color='0563C1', underline='single')
        rr += 1

        def put(head, text, italic=False):
            nonlocal rr
            h = rec.cell(rr, 1, head)
            h.font = f(9, True, color='546C78')
            h.alignment = Alignment(vertical='top')
            c = rec.cell(rr, 2, text)
            c.font = f(10, italic=italic)
            c.alignment = Alignment(wrap_text=True, vertical='top')
            lines = sum(len(s) // 150 + 1 for s in str(text).split('\n'))
            rec.row_dimensions[rr].height = min(400, 14 * max(1, lines))
            rr += 1

        put('Units', '\n'.join(x['units']) or '(none listed)')
        put('Email', x['email'] or '(none)')
        if x['url']:
            put('Profile page', x['url'])
        put('Profile text', x['profile'] or '(no profile text on the record)', italic=not x['profile'])
        put('Grant titles', '\n'.join(x['grants']) or '(none on the record)', italic=not x['grants'])
        pubs = x['pubs']
        head = 'Titles (%d)' % len(pubs)
        if not pubs:
            put(head, '(no publication titles on the record)', italic=True)
        for k, (t, y) in enumerate(pubs[:MAX_TITLES]):
            put(head if k == 0 else '', ('%s   ' % y if y else '') + t)
        if len(pubs) > MAX_TITLES:
            put('', '… and %d older titles, not shown' % (len(pubs) - MAX_TITLES), italic=True)
        for q in range(start + 2, rr):
            if (q - start) % 2 == 0:
                rec.cell(q, 2).fill = BAND
        rr += 2

        # ---- label row
        vals = [code, x['name'], '\n'.join(x['units'][:3]) + (' …' if len(x['units']) > 3 else ''),
                x['degrees'], len(pubs), '=HYPERLINK("#Records!A%d","Read record →")' % start]
        for j, v in enumerate(vals, 1):
            c = lab.cell(row, j, v)
            c.font = f(10)
            c.border = BOX
            c.alignment = Alignment(vertical='top', wrap_text=True)
        c = lab.cell(row, 6)
        c.font = f(10, color='0563C1', underline='single')
        for j in (7, 8, 9, 10):
            c = lab.cell(row, j)
            c.fill = INPUT
            c.border = BOX
            c.font = f(11 if j in (7, 8, 10) else 10, bold=j in (7, 10))
            c.alignment = Alignment(horizontal='center' if j != 9 else 'left', vertical='top', wrap_text=True)
        lab.cell(row, 5).alignment = Alignment(horizontal='center', vertical='top')
        lab.row_dimensions[row].height = 14 * max(1, min(3, len(x['units'])), min(4, len(x['degrees']) // 13 + 1))
    last = first + len(people) - 1
    dv_lab.add('G%d:G%d' % (first, last))
    dv_conf.add('H%d:H%d' % (first, last))
    dv_web.add('J%d:J%d' % (first, last))
    lab.freeze_panes = 'C%d' % (hr + 1)
    lab.auto_filter.ref = 'A%d:J%d' % (hr, last)

    # ------------------------------------------------------------ progress
    pr = wb.create_sheet('Progress')
    pr.sheet_view.showGridLines = False
    pr.column_dimensions['A'].width = 28
    pr.column_dimensions['B'].width = 12
    pr['A1'] = 'Progress'
    pr['A1'].font = f(14, True, color='1F3A4D')
    rng = lambda col: 'Label!%s%d:%s%d' % (col, first, col, last)
    rows = [('People to label', '=ROWS(%s)' % rng('A'), '0'),
            ('Labelled', '=COUNTA(%s)' % rng('G'), '0'),
            ('Left to do', '=B3-B4', '0'),
            ('Done', '=IF(B3=0,0,B4/B3)', '0%'),
            (None, None, None),
            ('W  wet', '=COUNTIF(%s,"W")' % rng('G'), '0'),
            ('D  dry', '=COUNTIF(%s,"D")' % rng('G'), '0'),
            ('H  hybrid', '=COUNTIF(%s,"H")' % rng('G'), '0'),
            ('U  can’t tell', '=COUNTIF(%s,"U")' % rng('G'), '0'),
            ('Labelled but Sure? blank', '=B4-COUNT(%s)' % rng('H'), '0')]
    for k, (a, b, fmt) in enumerate(rows, 3):
        if a is None:
            continue
        pr.cell(k, 1, a).font = f(10)
        c = pr.cell(k, 2, b)
        c.font = f(10, True)
        c.number_format = fmt
    pr.cell(15, 1, 'Counts update as you fill the Label sheet. The example row is not counted.').font = f(9, italic=True, color='546C78')

    for sh in (ws, lab, rec, pr):
        # printable: landscape, one page wide, header row repeated on the label sheet
        sh.page_setup.orientation = 'landscape'
        sh.page_setup.fitToWidth = 1
        sh.page_setup.fitToHeight = 0
        sh.sheet_properties.pageSetUpPr.fitToPage = True
    lab.print_title_rows = '%d:%d' % (hr, hr)
    wb.active = 0                       # Instructions, Label, Records, Progress
    wb.save(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('atlas')
    ap.add_argument('--split', default=os.path.join(os.path.dirname(os.path.abspath(__file__)), 'split.json'))
    ap.add_argument('--n', type=int, default=150)
    ap.add_argument('--out', required=True)
    a = ap.parse_args()
    D = W.load_data(a.atlas)
    test = sorted(json.load(open(a.split))['test'])
    rng = random.Random(SEED)
    pick = rng.sample(test, min(a.n, len(test)))        # random, and in random order
    people = [('H%03d' % (k + 1), i) for k, i in enumerate(pick)]
    os.makedirs(a.out, exist_ok=True)
    build(D, people, os.path.join(a.out, 'labelling_sheet.xlsx'))
    with open(os.path.join(a.out, 'key.csv'), 'w', newline='') as fh:
        w = csv.writer(fh)
        w.writerow(['code', 'index', 'name'])
        for code, i in people:
            w.writerow([code, i, D['pis'][i]['n']])
    print('wrote', len(people), 'people to', a.out)


if __name__ == '__main__':
    main()
