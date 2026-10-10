"""Score the human check of the wet/dry labels.

Reads one or more labelling workbooks filled by people (from
make_human_check.py) and answers three questions:

  1. Do the people agree with Claude's gold labels?   (is the answer key sound?)
  2. Do the people agree with each other?             (how hard is the call?)
  3. How often is the model right by the people's labels, next to how often it
     is right by Claude's labels on the same people?  (does 91.6% hold up?)

The model's calls are the locked-test calls in report/test_predictions.csv,
made by the version fitted on the training half only, so none of these people
was used to fit it.

  python ml/human_check.py ml/human_check/labelling_sheet_JS.xlsx [more.xlsx ...] \
      --key ml/human_check/key.csv --out ml/human_check/results

Writes results.md, disagreements.csv (every person where anyone disagrees, with
all the labels and notes side by side, for a joint review) and
human_labels.csv (every human label in long form).
"""
import argparse
import collections
import csv
import math
import os
import re
import sys

import numpy as np
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
SIDE = {0: None, 1: 'W', 4: 'W', 2: 'D', 5: 'D', 3: 'H'}
STATE_NAME = {0: 'unclassified', 1: 'wet', 4: 'leans wet', 3: 'hybrid', 5: 'leans dry', 2: 'dry'}
SEED = 20260926


def read_sheet(path):
    """{code: (label, sure, note, web)} from a filled workbook's Label sheet."""
    ws = load_workbook(path, data_only=True)['Label']
    head = None
    for r in range(1, 20):
        vals = [str(c.value or '').strip() for c in ws[r]]
        if 'Code' in vals and 'Label' in vals:
            head = r
            col = {v: j for j, v in enumerate(vals)}
            break
    if head is None:
        sys.exit(path + ': no header row with Code and Label on the Label sheet')
    out, bad = {}, []
    for row in ws.iter_rows(min_row=head + 1, values_only=True):
        code = str(row[col['Code']] or '').strip()
        if not code or code == 'EX':
            continue
        get = lambda k: row[col[k]] if k in col else None
        lab = str(get('Label') or '').strip().upper()
        sure = str(get('Sure?') or '').strip()
        sure = sure[:-2] if sure.endswith('.0') else sure
        web = str(get('After web check') or '').strip().upper()
        if lab and lab not in 'WDHU' or len(lab) > 1:
            bad.append('%s: label %r' % (code, lab))
            continue
        if web and (web not in 'WDHU' or len(web) > 1):
            bad.append('%s: web label %r' % (code, web))
            web = ''
        out[code] = (lab or None, sure if sure in ('1', '2') else None, str(get('Note') or '').strip(), web or None)
    if bad:
        print('%s: ignored %d unreadable cells: %s' % (path, len(bad), '; '.join(bad[:8])), file=sys.stderr)
    return out


def kappa(a, b, cats='WDHU'):
    """Cohen's kappa for two raters over the same items."""
    n = len(a)
    if not n:
        return float('nan')
    po = sum(x == y for x, y in zip(a, b)) / n
    ca, cb = collections.Counter(a), collections.Counter(b)
    pe = sum(ca[c] * cb[c] for c in cats) / (n * n)
    return (po - pe) / (1 - pe) if pe < 1 else float('nan')


def boot(fn, items, n=2000):
    """95% bootstrap interval of fn over resamples of items."""
    rng = np.random.default_rng(SEED)
    vals = []
    for _ in range(n):
        s = [items[j] for j in rng.integers(0, len(items), len(items))]
        v = fn(s)
        if not math.isnan(v):
            vals.append(v)
    return (float(np.percentile(vals, 2.5)), float(np.percentile(vals, 97.5))) if vals else (float('nan'),) * 2


def model_right(pairs):
    """pairs: [(reference label, model state)]. Share on the right side, scored
    on W/H/D references, with unclassified counted wrong (as in evaluate)."""
    sc = [(g, SIDE[s]) for g, s in pairs if g in ('W', 'H', 'D')]
    return sum(g == s for g, s in sc) / len(sc) if sc else float('nan')


def pct(x):
    return 'n/a' if x is None or (isinstance(x, float) and math.isnan(x)) else '%.1f%%' % (100 * x)


def table(rows_lab, cols_lab, pairs, rname, cname):
    t = collections.Counter(pairs)
    out = ['| %s \\ %s | %s |' % (rname, cname, ' | '.join(cols_lab)), '|' + '---|' * (len(cols_lab) + 1)]
    for r in rows_lab:
        out.append('| %s | %s |' % (r, ' | '.join(str(t.get((r, c), 0)) for c in cols_lab)))
    return '\n'.join(out)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('sheets', nargs='+', help='filled labelling workbooks, one per person')
    ap.add_argument('--names', help='comma-separated labeller names, in the order of the sheets (default: from file names)')
    ap.add_argument('--key', default=os.path.join(HERE, 'human_check', 'key.csv'))
    ap.add_argument('--labels', default=os.path.join(HERE, 'gold_labels.csv'))
    ap.add_argument('--test-predictions', default=os.path.join(HERE, 'report', 'test_predictions.csv'))
    ap.add_argument('--out', default=os.path.join(HERE, 'human_check', 'results'))
    a = ap.parse_args()

    key = {r['code']: (int(r['index']), r['name']) for r in csv.DictReader(open(a.key))}
    gold = {int(r['index']): r for r in csv.DictReader(open(a.labels))}
    tp = {int(r['index']): r for r in csv.DictReader(open(a.test_predictions))}
    for code, (i, name) in key.items():
        if i not in gold or gold[i]['name'] != name or gold[i]['split'] != 'test':
            sys.exit('key.csv %s points at record %d (%s), which is not that person in the test half of %s'
                     % (code, i, name, a.labels))
        if i not in tp:
            sys.exit('record %d (%s) is missing from %s' % (i, name, a.test_predictions))

    names = a.names.split(',') if a.names else [
        (re.sub(r'^labelling_sheet_?', '', os.path.splitext(os.path.basename(p))[0]) or 'rater%d' % (k + 1))
        for k, p in enumerate(a.sheets)]
    if len(names) != len(a.sheets) or len(set(names)) != len(names):
        sys.exit('need one distinct name per sheet: %s' % names)
    H = {}
    for nm, path in zip(names, a.sheets):
        got = read_sheet(path)
        unknown = set(got) - set(key)
        if unknown:
            sys.exit('%s has codes not in key.csv: %s' % (path, sorted(unknown)[:5]))
        H[nm] = {c: v for c, v in got.items() if v[0]}

    os.makedirs(a.out, exist_ok=True)
    claude = lambda c: gold[key[c][0]]['label']
    state = lambda c: int(tp[key[c][0]]['new_state'])
    L = ['# Human check of the wet/dry labels', '',
         'People labelled a random sample of the locked test set blind, under the same rubric Claude used. '
         'The model calls below are the locked-test calls, from the version fitted on the training half only.', '']

    # ---------------------------------------------------------- per labeller
    for nm, h in H.items():
        codes = sorted(h)
        n = len(codes)
        L += ['## %s: %d of %d labelled' % (nm, n, len(key)), '']
        if not n:
            L += ['Nothing labelled yet.', '']
            continue
        cnt = collections.Counter(h[c][0] for c in codes)
        L.append('Labels given: ' + ', '.join('%s %d' % (x, cnt[x]) for x in 'WDHU') +
                 '; marked sure on %d.' % sum(h[c][1] == '1' for c in codes))
        L.append('')
        # 1. agreement with Claude
        pairs = [(h[c][0], claude(c)) for c in codes]
        agree = lambda s: sum(x == y for x, y in s) / len(s)
        wet = lambda s: sum((x == 'W') == (y == 'W') for x, y in s) / len(s)
        lo, hi = boot(agree, pairs)
        wlo, whi = boot(wet, pairs)
        sure = [(h[c][0], claude(c)) for c in codes if h[c][1] == '1' and gold[key[c][0]]['conf'] == '1']
        L += ['**Agreement with Claude\'s labels**', '',
              '| | agreement | 95% CI | Cohen\'s kappa |', '|---|---|---|---|',
              '| same label (W/D/H/U) | %s | %s-%s | %.2f |' % (pct(agree(pairs)), pct(lo), pct(hi),
                                                              kappa([x for x, _ in pairs], [y for _, y in pairs])),
              '| wet vs not wet | %s | %s-%s | %.2f |' % (pct(wet(pairs)), pct(wlo), pct(whi),
                                                        kappa(['W' if x == 'W' else 'N' for x, _ in pairs],
                                                              ['W' if y == 'W' else 'N' for _, y in pairs], 'WN')),
              '| same label, both sure | %s | | (n=%d) |' % (pct(agree(sure)) if sure else 'n/a', len(sure)), '',
              table('WDHU', list('WDHU'), pairs, nm, 'Claude'), '']
        # 3. the model, judged by this person and by Claude on the same people
        mh = [(h[c][0], state(c)) for c in codes]
        mc = [(claude(c), state(c)) for c in codes]
        both = list(zip(mh, mc))
        d = lambda s: model_right([x for x, _ in s]) - model_right([y for _, y in s])
        dlo, dhi = boot(d, both)
        L += ['**The model, by these labels**', '',
              '| judged by | model on the right side | 95% CI |', '|---|---|---|',
              '| %s | %s | %s-%s |' % ((nm, pct(model_right(mh))) + tuple(pct(v) for v in boot(model_right, mh))),
              '| Claude (same people) | %s | %s-%s |' % ((pct(model_right(mc)),) + tuple(pct(v) for v in boot(model_right, mc))),
              '', 'Difference (%s minus Claude): %+.1f to %+.1f points. An interval that holds zero means the '
              'published accuracy stands on these labels.' % (nm, 100 * dlo, 100 * dhi), '']
        web = [(c, h[c][3]) for c in codes if h[c][3] and h[c][3] != h[c][0]]
        if any(h[c][3] for c in codes):
            wref = [((h[c][3] or h[c][0]), state(c)) for c in codes]
            L += ['After checking the lab\'s own web page or what they knew of it, %s changed %d label%s. '
                  'By those final labels the model is right for %s.' % (nm, len(web), '' if len(web) == 1 else 's',
                                                                        pct(model_right(wref))), '']

    # ------------------------------------------------------- between people
    names_done = [nm for nm in H if H[nm]]
    if len(names_done) >= 2:
        L += ['## Between the people', '']
        for x in range(len(names_done)):
            for y in range(x + 1, len(names_done)):
                A, B = H[names_done[x]], H[names_done[y]]
                cs = sorted(set(A) & set(B))
                if not cs:
                    continue
                aa, bb = [A[c][0] for c in cs], [B[c][0] for c in cs]
                L.append('- %s and %s, %d people in common: same label on %s (kappa %.2f); on wet vs not '
                         'wet, %s.' % (names_done[x], names_done[y], len(cs),
                                       pct(sum(p == q for p, q in zip(aa, bb)) / len(cs)), kappa(aa, bb),
                                       pct(sum((p == 'W') == (q == 'W') for p, q in zip(aa, bb)) / len(cs))))
        cons = {}
        for c in key:
            got = [H[nm][c][0] for nm in names_done if c in H[nm]]
            if len(got) >= 2 and len(set(got)) == 1:
                cons[c] = got[0]
        if cons:
            pairs = [(cons[c], claude(c)) for c in cons]
            L += ['', 'Where every person agreed (%d people), Claude gave the same label on %s, and the model is '
                  'right for %s of them.' % (len(cons), pct(sum(p == q for p, q in pairs) / len(pairs)),
                                             pct(model_right([(cons[c], state(c)) for c in cons])))]
        L.append('')

    # ----------------------------------------------------------- the files
    L += ['## What to do next', '',
          '`disagreements.csv` lists every person on whom anyone differs from Claude, with every label and note '
          'side by side. Go through it together; where the people are right, change that row of '
          '`ml/gold_labels.csv` (label, conf, and a note saying it was corrected by the human check) and rebuild '
          'with `scripts/build_atlas.sh`. The model retrains on the corrected labels and every figure on the page '
          'is recomputed.', '']
    open(os.path.join(a.out, 'results.md'), 'w').write('\n'.join(L) + '\n')

    with open(os.path.join(a.out, 'human_labels.csv'), 'w', newline='') as fh:
        w = csv.writer(fh)
        w.writerow(['code', 'index', 'name', 'labeller', 'label', 'sure', 'note', 'after_web_check'])
        for nm, h in H.items():
            for c in sorted(h):
                w.writerow([c, key[c][0], key[c][1], nm] + list(h[c]))
    with open(os.path.join(a.out, 'disagreements.csv'), 'w', newline='') as fh:
        w = csv.writer(fh)
        w.writerow(['code', 'index', 'name', 'claude_label', 'claude_note'] +
                   [x for nm in H for x in (nm + '_label', nm + '_note')] +
                   ['model_call', 'p_wet', 'model_confidence'])
        for c in sorted(key):
            labs = {nm: H[nm][c] for nm in H if c in H[nm]}
            if not labs or all(v[0] == claude(c) for v in labs.values()):
                continue
            i = key[c][0]
            w.writerow([c, i, key[c][1], claude(c), gold[i]['note']] +
                       [x for nm in H for x in ((labs[nm][0], labs[nm][2]) if nm in labs else ('', ''))] +
                       [STATE_NAME[state(c)], tp[i]['p_wet'], tp[i].get('confidence', '')])
    print(open(os.path.join(a.out, 'results.md')).read())


if __name__ == '__main__':
    main()
